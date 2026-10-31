"""End-to-end regression coverage for routing precedence and tool-plan contracts."""
from __future__ import annotations

import os
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from pptx import Presentation

from eva.automation.service import AutomationService
from eva.automation.task_planner import TaskPlanner, _tool_rows
from eva.automation.content_contracts import validate_generated_content
from eva.core.intent_understanding import analyze_intent
from eva.core.task_router import TaskRouter


class FixedPresentationProvider:
    def generate_text(self, system_prompt, user_prompt, *, timeout=45, num_predict=1200):
        return ("Renewable Energy\n"
                "Slide 1: Overview\n- Renewable energy comes from replenished sources\n- Examples include solar and wind\n"
                "---SLIDE---\n"
                "Slide 2: Solar Power\n- Solar panels convert sunlight into electricity\n- Output depends on sunlight and system design\n"
                "---SLIDE---\n"
                "Slide 3: Wind Power\n- Turbines convert moving air into electricity\n- Siting affects output and local impacts\n"
                "---SLIDE---\n"
                "Slide 4: Benefits and Challenges\n- Operational emissions are generally low\n- Storage and grid integration can be challenges\n"
                "---SLIDE---\n"
                "Slide 5: Conclusion\n- A diverse energy mix improves resilience\n- Planning should reflect local resources")


class RoutingPipelineHardeningTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix="eva-routing-pipeline-")
        self.root = Path(self.temp.name).resolve()
        for folder in ("Desktop", "Documents", "Downloads"):
            (self.root / folder).mkdir()
        self.service = AutomationService(
            allowed_roots=[self.root], audit_path=None,
            workflows_dir=self.root / "workflows",
            capability_gaps_path=self.root / "state" / "gaps.json",
            browser_headless=True,
        )
        self.planner = TaskPlanner(self.service, None, self.service.capability_gaps)

    def tearDown(self):
        self.service.close()
        self.temp.cleanup()

    def test_direct_url_plus_find_text_is_an_ordered_browser_plan(self):
        prompt = "Open https://example.com and find the word Example on the page"
        self.assertEqual(analyze_intent(prompt).intent, "web_browsing")
        plan = self.planner._deterministic_plan(prompt)
        self.assertIsNotNone(plan)
        self.assertEqual([step["tool"] for step in plan["steps"]], ["browser.navigate", "browser.find_text"])
        self.assertEqual(plan["steps"][0]["args"]["url"], "https://example.com")
        self.assertEqual(plan["steps"][1]["args"]["query"], "Example")
        self.assertEqual(plan["steps"][1]["depends_on"], ["open_page"])

    def test_search_then_download_is_not_silently_reduced_to_search_only(self):
        prompt = "Search the web for the latest Nvidia driver and download it"
        self.assertIsNone(self.planner._deterministic_plan(prompt))
        names = {row["name"] for row in _tool_rows(self.service, prompt)}
        self.assertIn("browser.snapshot", names)
        self.assertIn("browser.download_link", names)

    def test_direct_public_url_download_is_recognized_as_download_intent(self):
        from eva.automation.task_planner import _extract_web_browse_request
        request = _extract_web_browse_request("Download https://example.com/installer.exe")
        self.assertIsNotNone(request)
        self.assertTrue(request["wants_download"])
        self.assertEqual(request["url"], "https://example.com/installer.exe")
        self.assertIsNone(self.planner._deterministic_plan("Download https://example.com/installer.exe"))

    def test_official_documentation_request_becomes_api_free_browse_and_summary(self):
        prompt = "Find the official Python documentation and summarize the installation instructions"
        plan = self.planner._deterministic_plan(prompt)
        self.assertIsNotNone(plan)
        self.assertEqual(plan["steps"][0]["tool"], "web.browse")
        self.assertIn("Python", plan["steps"][0]["args"]["query"])
        self.assertIn("installation instructions", plan["steps"][0]["args"]["query"])
        self.assertEqual(plan["steps"][1]["tool"], "content.generate_text")
        self.assertIn("${steps.browse.data.research_text}", plan["steps"][1]["args"]["prompt"])

    def test_gmail_spam_task_is_not_a_local_file_resolution(self):
        prompt = "Open Gmail and show my spam folder"
        plan = self.planner._deterministic_plan(prompt)
        self.assertIsNotNone(plan)
        self.assertEqual([step["tool"] for step in plan["steps"]], [
            "browser.navigate", "browser.snapshot", "browser.click_by_role"
        ])
        self.assertEqual(plan["steps"][2]["args"]["name"], "Spam")
        self.assertIn("isolated browser session", plan["steps"][0]["purpose"])
        rows = _tool_rows(self.service, prompt)
        names = {row["name"] for row in rows}
        self.assertTrue({"browser.navigate", "browser.snapshot", "browser.click_by_role"} <= names)

    def test_latest_pdf_uses_modification_time_not_folder_listing(self):
        older = self.root / "Downloads" / "older.pdf"
        newer = self.root / "Downloads" / "newer.pdf"
        older.write_bytes(b"older")
        newer.write_bytes(b"newer")
        os.utime(older, (1_000_000_000, 1_000_000_000))
        os.utime(newer, (1_100_000_000, 1_100_000_000))

        # A real home Downloads directory can exist but be outside this service's
        # explicitly approved root. It must not override the in-root test folder.
        with tempfile.TemporaryDirectory(prefix="eva-unapproved-home-") as foreign_home_dir:
            foreign_home = Path(foreign_home_dir)
            foreign_downloads = foreign_home / "Downloads"
            foreign_downloads.mkdir()
            outside_newest = foreign_downloads / "outside-newest.pdf"
            outside_newest.write_bytes(b"outside")
            os.utime(outside_newest, (1_200_000_000, 1_200_000_000))
            with patch("eva.automation.service.Path.home", return_value=foreign_home):
                result = self.service.handle_text("Find the most recent PDF in my Downloads folder")

        self.assertTrue(result.get("ok"), result)
        self.assertEqual(result.get("tool"), "files.find_latest")
        self.assertEqual(Path(result["data"]["path"]), newer)
        self.assertNotEqual(Path(result["data"]["path"]), outside_newest)
        self.assertTrue(result.get("verified"))

    def test_readme_summary_reads_source_before_generating_and_saving(self):
        prompt = "Read README, explain, and save a summary as README.md"
        plan = self.planner._deterministic_plan(prompt)
        self.assertIsNotNone(plan)
        self.assertEqual([step["tool"] for step in plan["steps"]], [
            "files.resolve_by_name", "files.read_document", "content.generate_text", "files.create_unique_text_file"
        ])
        self.assertEqual(plan["steps"][0]["args"]["filename"], "README.md")
        self.assertIn("${steps.read_readme.data.text}", plan["steps"][2]["args"]["prompt"])
        self.assertEqual(plan["steps"][3]["args"]["filename"], "README.md")
        self.assertEqual(plan["steps"][3]["depends_on"], ["summarize_readme"])

    def test_readme_summary_workflow_executes_without_overwriting_source(self):
        class SummaryProvider:
            def generate_text(self, system_prompt, user_prompt, *, timeout=45, num_predict=1200):
                self.last_prompt = user_prompt
                return "- The README describes the project.\n- It documents local tools."
        provider = SummaryProvider()
        source = self.root / "README.md"
        source_text = "# EVA Demo\n\nThis project documents local tools."
        source.write_text(source_text, encoding="utf-8")
        with patch.object(Path, "cwd", return_value=self.root):
            result = self.service.handle_text(
                "Read README, explain, and save a summary as README.md", provider=provider
            )
        self.assertTrue(result.get("ok"), result)
        self.assertTrue(result.get("verified"), result)
        output = self.root / "Documents" / "EVA-created-files" / "README.md"
        self.assertTrue(output.is_file())
        self.assertEqual(output.read_text(encoding="utf-8"), "- The README describes the project.\n- It documents local tools.")
        self.assertEqual(source.read_text(encoding="utf-8"), source_text)
        self.assertIn(source_text, provider.last_prompt)

    def test_powerpoint_plan_creates_real_multislide_file(self):
        prompt = "Create a PowerPoint presentation about renewable energy with 5 slides"
        plan = self.planner._deterministic_plan(prompt)
        self.assertIsNotNone(plan)
        self.assertEqual([step["tool"] for step in plan["steps"]], ["content.generate_text", "files.create_pptx"])
        self.assertIn("exactly 5 slides", plan["steps"][0]["args"]["prompt"])
        self.assertEqual(plan["steps"][0]["args"]["output_kind"], "presentation")
        self.assertEqual(plan["steps"][0]["args"]["expected_count"], 5)
        self.assertEqual(plan["steps"][1]["args"]["path"].endswith(".pptx"), True)
        good_outline = "Renewable Energy\nSlide 1: Overview\nClean sources\n---SLIDE---\nSlide 2: Solar\nSunlight"
        self.assertIsNone(validate_generated_content(good_outline, "presentation", "renewable energy", 2))
        bad_outline = "Renewable Energy\nSlide 1: Overview\nClean sources\n---SLIDE---\nSlide 2: Solar\nSunlight\n---SLIDE---\nSlide 3: Wind\nTurbines"
        self.assertIn("exactly 2", validate_generated_content(bad_outline, "presentation", "renewable energy", 2) or "")

        path = self.root / "Documents" / "presentation-smoke.pptx"
        content = (
            "Renewable Energy\n"
            "Slide 1: Overview\n- Clean energy sources\n- Lower emissions\n"
            "---SLIDE---\n"
            "Slide 2: Solar\n- Sunlight generates electricity\n- Panels need suitable sites\n"
            "---SLIDE---\n"
            "Slide 3: Wind\n- Turbines convert wind energy\n- Output depends on wind conditions"
        )
        result = self.service.execute_tool("files.create_pptx", {"path": str(path), "content": content})
        self.assertTrue(result.ok, result.as_dict())
        self.assertTrue(result.verified)
        reopened = Presentation(str(path))
        self.assertEqual(len(reopened.slides), 3)
        slide_text = [" ".join(shape.text for shape in slide.shapes if getattr(shape, "has_text_frame", False)) for slide in reopened.slides]
        self.assertIn("Overview", slide_text[0])
        self.assertIn("Solar", slide_text[1])
        self.assertIn("Wind", slide_text[2])

    def test_powerpoint_workflow_executes_end_to_end_with_local_provider(self):
        prompt = "Create a PowerPoint presentation about renewable energy with 5 slides"
        result = self.service.handle_text(prompt, provider=FixedPresentationProvider())
        self.assertTrue(result.get("ok"), result)
        self.assertTrue(result.get("verified"), result)
        create_step = next(step for step in result["execution"]["steps"] if step["tool"] == "files.create_pptx")
        path = Path(create_step["result"]["data"]["path"])
        reopened = Presentation(str(path))
        self.assertEqual(len(reopened.slides), 5)
        all_text = "\n".join(shape.text for slide in reopened.slides for shape in slide.shapes if getattr(shape, "has_text_frame", False))
        self.assertIn("Solar Power", all_text)
        self.assertIn("Conclusion", all_text)

    def test_mixed_office_and_code_request_is_delegated_to_the_orchestrator(self):
        plan = self.planner._deterministic_plan(
            "Create an Excel expense tracker and a Python script that reads it"
        )
        self.assertIsNone(plan, "A single-output deterministic shortcut must not swallow a compound cross-domain request")

    def test_multi_intent_schema_shortlist_keeps_web_and_document_creation_tools(self):
        rows = _tool_rows(self.service,
            "Research the official Python docs, summarize the installation steps, and save a Word document called setup.docx")
        names = {row["name"] for row in rows}
        self.assertTrue({"web.browse", "content.generate_text", "files.create_docx"} <= names, names)

    def test_tool_schema_shortlist_keeps_browser_navigation_and_text_finding(self):
        rows = _tool_rows(self.service, "Open https://example.com and find the word Example on the page")
        names = {row["name"] for row in rows}
        self.assertTrue({"browser.navigate", "browser.find_text"} <= names)

    def test_gmail_folder_request_uses_default_browser_profile_and_does_not_claim_verified_login(self):
        from eva.automation.engine import ToolResult
        expected_url = "https://mail.google.com/mail/u/0/#spam"
        with patch.object(self.service, "execute_tool", return_value=ToolResult(
            True, "desktop.open_url", {"url": expected_url, "launch_requested": True},
            "Default browser launch requested; page rendering not verified.", verified=False
        )) as execute:
            result = self.service.handle_text("Open Gmail and show my spam folder")
        self.assertEqual(execute.call_args.args[0], "desktop.open_url")
        self.assertEqual(execute.call_args.args[1]["url"], expected_url)
        self.assertEqual(result["action"], "gmail_folder_open")
        self.assertIn("does not verify Gmail sign-in", result["note"])

    def test_pathless_project_inspection_uses_approved_current_project(self):
        (self.root / "main.py").write_text("def main():\n    return 0\n", encoding="utf-8")
        (self.root / "src").mkdir()
        (self.root / "src" / "helper.py").write_text("VALUE = 1\n", encoding="utf-8")
        with patch.object(Path, "cwd", return_value=self.root):
            result = self.service.handle_text(
                "Inspect this project and tell me its folder structure and main entry point"
            )
        self.assertTrue(result.get("ok"), result)
        self.assertEqual(result.get("action"), "project_inspection")
        self.assertIn("main.py", result["data"]["tree"])
        self.assertTrue(result.get("verified"))

    def test_mixed_office_and_code_request_without_a_local_planner_fails_safely(self):
        result = self.service.handle_text(
            "Create an Excel expense tracker and a Python script that reads it"
        )
        self.assertFalse(result.get("ok"), result)
        self.assertNotEqual(result.get("action"), "capability_gap", "Do not report a fake missing orchestrator while registered tools exist")
        self.assertFalse(list((self.root / "Documents" / "EVA-created-files").glob("*.xlsx")))

    def test_simple_math_and_local_file_intents_do_not_regress(self):
        router = TaskRouter()
        self.assertEqual(router.route("What is 18 percent of 450?").handler, "math_engine")
        self.assertEqual(analyze_intent("Find the latest report and save the summary as final.txt").intent,
                         "file_inspection")
        self.assertEqual(router.route("What is today's date?").handler, "system_time")

    def test_audit_reproduced_routing_failures_now_route_to_automation(self):
        router = TaskRouter()
        queries = [
            "Copy the latest PDF from Downloads.",
            "Move older PDFs.",
            "Add a bottle to my Amazon cart.",
            "Lock my screen.",
            "Delete all duplicate files in Downloads.",
        ]
        for q in queries:
            decision = router.route(q)
            self.assertEqual(decision.handler, "desktop_automation", f"Query {q!r} routed to {decision.handler}")
            self.assertEqual(decision.category, "DESKTOP_AUTOMATION", f"Query {q!r} category {decision.category}")


if __name__ == "__main__":
    unittest.main()
